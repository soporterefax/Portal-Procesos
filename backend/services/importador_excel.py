from __future__ import annotations

import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from extensions import db
from models import Area, Documento, Proceso


RUTA_EXCEL_DEFAULT = (
    Path(__file__).resolve().parent.parent
    / "data"
    / "Plantilla_Procesos_Refax.xlsx"
)

HOJAS_REQUERIDAS = {"AREAS", "PROCESOS", "DOCUMENTOS"}

COLUMNAS_AREAS = {
    "CODIGO_AREA",
    "NOMBRE_AREA",
    "RESPONSABLE_AREA",
    "NOMBRE_PERSONAL",
    "DESCRIPCION",
}

COLUMNAS_PROCESOS = {
    "CODIGO",
    "AREA",
    "TIPO",
    "NOMBRE_PROCESO",
    "RESPONSABLE",
    "PERSONA_RESPONSABLE",
    "OBJETIVO",
    "ES_CRITICO",
    "AREAS_RELACIONADAS",
    "FECHA_ACTUALIZACION",
}

COLUMNAS_DOCUMENTOS = {
    "CODIGO_PROCESO",
    "TIPO_DOCUMENTO",
    "NOMBRE_DOCUMENTO",
    "VERSION",
    "FECHA_ACTUALIZACION",
    "ENLACE_DOC",
    "ENLACE_FLUJ",
    "ENLACE_FLUJ1",
    "ENLACE_FLUJ2",
    "ENLACE_FLUJ3",
}


def limpiar_texto(valor):
    if valor is None:
        return None
    texto = str(valor).strip()
    if not texto or texto.lower() in {"nan", "none"}:
        return None
    return texto


def normalizar_codigo(valor):
    texto = limpiar_texto(valor)
    return texto.upper() if texto else None


def normalizar_nombre(valor):
    texto = limpiar_texto(valor)
    if not texto:
        return None
    texto = " ".join(texto.upper().split())
    texto = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in texto if not unicodedata.combining(c)) or None


def convertir_booleano(valor):
    texto = limpiar_texto(valor)
    if not texto:
        return False
    return texto.lower() in {"sí", "si", "s", "1", "true", "verdadero", "x"}


def convertir_fecha(valor):
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = limpiar_texto(valor)
    if not texto:
        return None
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            pass
    return None


def filas_como_dict(worksheet):
    filas = worksheet.iter_rows(values_only=True)
    try:
        encabezados = [str(v).strip().upper() if v is not None else "" for v in next(filas)]
    except StopIteration:
        return [], set()

    registros = []
    for numero_fila, valores in enumerate(filas, start=2):
        if not any(v is not None and str(v).strip() for v in valores):
            continue
        fila = dict(zip(encabezados, valores))
        fila["__FILA__"] = numero_fila
        registros.append(fila)
    return registros, set(encabezados)


def validar_columnas(columnas_actuales, columnas_requeridas, hoja):
    faltantes = columnas_requeridas - columnas_actuales
    if faltantes:
        raise ValueError(
            f"La hoja {hoja} no contiene estas columnas: {', '.join(sorted(faltantes))}"
        )


def importar_areas(filas):
    creadas = actualizadas = omitidas = 0
    for fila in filas:
        codigo = normalizar_codigo(fila.get("CODIGO_AREA"))
        nombre = limpiar_texto(fila.get("NOMBRE_AREA"))
        if not codigo or not nombre:
            omitidas += 1
            continue

        area = Area.query.filter(db.func.upper(db.func.trim(Area.codigo)) == codigo).first()
        if area:
            actualizadas += 1
        else:
            area = Area(codigo=codigo, nombre=nombre)
            db.session.add(area)
            creadas += 1

        area.codigo = codigo
        area.nombre = nombre
        area.responsable_area = limpiar_texto(fila.get("RESPONSABLE_AREA"))
        area.nombre_personal = limpiar_texto(fila.get("NOMBRE_PERSONAL"))
        area.descripcion = limpiar_texto(fila.get("DESCRIPCION"))

    db.session.flush()
    return creadas, actualizadas, omitidas


def importar_procesos(filas):
    creados = actualizados = omitidos = 0

    mapa_areas = {}
    for area in Area.query.all():
        for valor in (area.codigo, area.nombre):
            clave = normalizar_nombre(valor)
            if clave:
                mapa_areas[clave] = area

    for fila in filas:
        codigo = normalizar_codigo(fila.get("CODIGO"))
        valor_area = limpiar_texto(fila.get("AREA"))
        nombre = limpiar_texto(fila.get("NOMBRE_PROCESO"))

        if not codigo or not valor_area or not nombre:
            omitidos += 1
            continue

        area = mapa_areas.get(normalizar_nombre(valor_area))
        if not area:
            print(
                f"Proceso omitido, área no encontrada (fila {fila.get('__FILA__')}): "
                f"{codigo} | {valor_area} | {nombre}"
            )
            omitidos += 1
            continue

        proceso = Proceso.query.filter(
            db.func.upper(db.func.trim(Proceso.codigo)) == codigo
        ).first()

        if proceso:
            actualizados += 1
        else:
            proceso = Proceso(codigo=codigo, area_id=area.id, nombre=nombre)
            db.session.add(proceso)
            creados += 1

        proceso.codigo = codigo
        proceso.area_id = area.id
        proceso.tipo = limpiar_texto(fila.get("TIPO"))
        proceso.nombre = nombre
        proceso.responsable = limpiar_texto(fila.get("RESPONSABLE"))
        proceso.persona_responsable = limpiar_texto(fila.get("PERSONA_RESPONSABLE"))
        proceso.objetivo = limpiar_texto(fila.get("OBJETIVO"))
        proceso.es_critico = convertir_booleano(fila.get("ES_CRITICO"))
        proceso.areas_relacionadas = limpiar_texto(fila.get("AREAS_RELACIONADAS"))
        proceso.fecha_actualizacion = convertir_fecha(fila.get("FECHA_ACTUALIZACION"))

    db.session.flush()
    return creados, actualizados, omitidos


def importar_documentos(filas):
    creados = actualizados = omitidos = 0

    for fila in filas:
        codigo_proceso = normalizar_codigo(fila.get("CODIGO_PROCESO"))
        tipo = limpiar_texto(fila.get("TIPO_DOCUMENTO"))
        nombre = limpiar_texto(fila.get("NOMBRE_DOCUMENTO"))

        if not codigo_proceso or not tipo or not nombre:
            omitidos += 1
            continue

        proceso = Proceso.query.filter(
            db.func.upper(db.func.trim(Proceso.codigo)) == codigo_proceso
        ).first()
        if not proceso:
            print(
                f"Documento omitido, proceso no encontrado (fila {fila.get('__FILA__')}): "
                f"{codigo_proceso} | {nombre}"
            )
            omitidos += 1
            continue

        documento = Documento.query.filter(
            Documento.proceso_id == proceso.id,
            db.func.lower(db.func.trim(Documento.tipo)) == tipo.lower(),
            db.func.lower(db.func.trim(Documento.nombre)) == nombre.lower(),
        ).first()

        if documento:
            actualizados += 1
        else:
            documento = Documento(proceso_id=proceso.id, tipo=tipo, nombre=nombre)
            db.session.add(documento)
            creados += 1

        documento.proceso_id = proceso.id
        documento.tipo = tipo
        documento.nombre = nombre
        documento.version = limpiar_texto(fila.get("VERSION"))
        documento.fecha_actualizacion = convertir_fecha(fila.get("FECHA_ACTUALIZACION"))
        documento.enlace_doc = limpiar_texto(fila.get("ENLACE_DOC"))
        documento.enlace_fluj = limpiar_texto(fila.get("ENLACE_FLUJ"))
        documento.enlace_fluj1 = limpiar_texto(fila.get("ENLACE_FLUJ1"))
        documento.enlace_fluj2 = limpiar_texto(fila.get("ENLACE_FLUJ2"))
        documento.enlace_fluj3 = limpiar_texto(fila.get("ENLACE_FLUJ3"))

    db.session.flush()
    return creados, actualizados, omitidos


def importar_excel(ruta_excel=None):
    ruta = Path(ruta_excel) if ruta_excel else RUTA_EXCEL_DEFAULT
    if not ruta.exists():
        raise FileNotFoundError(f"No se encontró el Excel: {ruta}")

    workbook = load_workbook(ruta, read_only=True, data_only=True)
    mapa_hojas = {nombre.strip().upper(): nombre for nombre in workbook.sheetnames}
    faltantes = HOJAS_REQUERIDAS - set(mapa_hojas)
    if faltantes:
        raise ValueError(f"Faltan las hojas: {', '.join(sorted(faltantes))}")

    areas, cols_areas = filas_como_dict(workbook[mapa_hojas["AREAS"]])
    procesos, cols_procesos = filas_como_dict(workbook[mapa_hojas["PROCESOS"]])
    documentos, cols_documentos = filas_como_dict(workbook[mapa_hojas["DOCUMENTOS"]])

    validar_columnas(cols_areas, COLUMNAS_AREAS, "AREAS")
    validar_columnas(cols_procesos, COLUMNAS_PROCESOS, "PROCESOS")
    validar_columnas(cols_documentos, COLUMNAS_DOCUMENTOS, "DOCUMENTOS")

    try:
        resultado_areas = importar_areas(areas)
        resultado_procesos = importar_procesos(procesos)
        resultado_documentos = importar_documentos(documentos)
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    finally:
        workbook.close()

    return {
        "ruta": str(ruta),
        "areas": {
            "creadas": resultado_areas[0],
            "actualizadas": resultado_areas[1],
            "omitidas": resultado_areas[2],
        },
        "procesos": {
            "creados": resultado_procesos[0],
            "actualizados": resultado_procesos[1],
            "omitidos": resultado_procesos[2],
        },
        "documentos": {
            "creados": resultado_documentos[0],
            "actualizados": resultado_documentos[1],
            "omitidos": resultado_documentos[2],
        },
    }
